"""docx 어댑터 — **순수 파이썬이고 덮어쓰기가 된다** (plan/24 T-0151 · ADR-0060).

hwp 와 갈리는 점이 둘이다.

- **JVM 이 없다.** 어디서든 돈다.
- **덮어쓸 수 있다.** XML 이라 갈아 끼워도 무너지지 않는다(실측에서 4칸을 덮어쓰고
  확인했다). 안내문이 든 칸을 지우고 값을 넣을 수 있어서 **hwp 에서 막힌 자리가
  여기서는 열린다** — 창업진흥원이 같은 양식을 hwp·docx 두 벌로 붙이는 것이
  그래서 기회가 된다.

## `cell.text = value` 를 쓰지 않는다

그 대입은 셀의 문단과 런을 전부 버리고 새로 만들어서 **글꼴·크기·정렬이 통째로
날아간다**(ADR-0060 결정 2). 첫 런의 글자만 바꾸고 나머지 런을 비우면 그 자리의
서식이 남는다.
"""

from __future__ import annotations

from finnect.forms.base import Anchor, Cell, FormsUnavailable


class DocxForm:
    """docx 파일 하나. **빈 칸이 아니어도 채운다.**"""

    ext = "docx"

    def __init__(self, path: str) -> None:
        try:
            import docx
        except ImportError as e:
            raise FormsUnavailable(
                "docx 를 열려면 `uv sync --group forms` 가 필요해요"
            ) from e

        self._path = path
        self._doc = docx.Document(path)

    def _tables(self):  # noqa: ANN202
        """최상위 표만. **중첩 표를 안 편다.**

        `python-docx` 의 `doc.tables` 가 최상위만 주고, 중첩까지 펴면 hwp 쪽 순회와
        번호가 어긋난다 — 좌표는 포맷을 건너 같은 뜻이어야 한다. 중첩 표에 든 칸은
        좌표를 못 얻고, 그 칸은 안 채워진다(`anchor` 가 NULL 인 것과 같은 결과다).
        """
        return list(self._doc.tables)

    def cells(self) -> list[Cell]:
        out: list[Cell] = []
        for ti, table in enumerate(self._tables()):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    # **구역이 늘 0 이다.** docx 는 본문이 하나라 표 번호가 문서
                    # 전체에서 이어진다(hwp 는 구역 안에서만 이어진다).
                    out.append(Cell(Anchor(0, ti, ri, ci), cell.text.strip()))
        return out

    def _cell_at(self, a: Anchor):  # noqa: ANN202
        if a.section != 0:
            return None
        tables = self._tables()
        if a.table >= len(tables):
            return None
        rows = tables[a.table].rows
        if a.row >= len(rows):
            return None
        cells = rows[a.row].cells
        return cells[a.col] if a.col < len(cells) else None

    def fill(self, anchor: Anchor, value: str) -> bool:
        """**덮어쓴다.** 안내문이 들어 있어도 지우고 값을 넣는다.

        서식을 지키려고 **첫 런만 갈고 나머지를 비운다.** 문단이 여럿이면 첫
        문단에 값을 넣고 나머지 문단의 런을 비운다 — 문단 자체를 지우면 셀의
        높이가 바뀌어 표가 다시 흐른다.
        """
        cell = self._cell_at(anchor)
        if cell is None or not cell.paragraphs:
            return False

        head, *rest = cell.paragraphs
        if head.runs:
            head.runs[0].text = value
            for r in head.runs[1:]:
                r.text = ""
        else:
            head.add_run(value)
        for para in rest:
            for r in para.runs:
                r.text = ""
        return True

    def save(self, path: str) -> None:
        self._doc.save(path)
