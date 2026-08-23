"""DOCX 양식의 표 셀을 라벨로 찾아 채운다. JVM 이 필요 없다.

`cell.text = value` 를 쓰지 않는다. python-docx 의 그 대입은 셀의 문단과 런을 전부
버리고 새로 만들어서 글꼴·크기·정렬이 통째로 날아간다. 첫 런의 글자만 바꾸고
나머지 런을 비우면 그 자리의 서식이 남는다.

실행:
    uv run --with python-docx python spike/docx_fill.py
"""

from __future__ import annotations

from pathlib import Path

from form_values import norm, 값, 빈칸표시

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "신청서"
OUT = ROOT / "data" / "신청서-채움"


def 모든표(container):
    """중첩 표까지 훑는다."""
    for t in container.tables:
        yield t
        for row in t.rows:
            for cell in row.cells:
                yield from 모든표(cell)


def set_cell(cell, value: str) -> None:
    """첫 문단의 글자만 갈아 끼운다. 문단과 런의 서식은 남긴다."""
    para = cell.paragraphs[0]
    if para.runs:
        para.runs[0].text = value
        for r in para.runs[1:]:
            r.text = ""
    else:
        para.add_run(value)


def fill(doc) -> tuple[list[tuple[str, str]], list[tuple[str, str]]]:
    새로, 덮어씀 = [], []
    본적 = set()
    for t in 모든표(doc):
        for row in t.rows:
            cells = row.cells
            for ci, cell in enumerate(cells):
                라벨 = norm(cell.text)
                if 라벨 not in 값:
                    continue
                for cj in range(ci + 1, len(cells)):
                    옆 = cells[cj]
                    본문 = (옆.text or "").strip()
                    if norm(본문) == 라벨:  # 병합 셀이라 라벨이 반복된다
                        continue
                    if not 빈칸표시.match(본문):
                        break
                    키 = (id(옆._tc), 라벨)
                    if 키 in 본적:
                        break
                    본적.add(키)
                    set_cell(옆, 값[라벨])
                    (덮어씀 if 본문 else 새로).append((라벨, 본문 or 값[라벨]))
                    break
    return 새로, 덮어씀


def 통계(doc) -> tuple[int, int, int]:
    ts = list(모든표(doc))
    셀 = sum(len(r.cells) for t in ts for r in t.rows)
    return len(ts), 셀, len(doc.paragraphs)


def main() -> None:
    from docx import Document

    OUT.mkdir(parents=True, exist_ok=True)
    files = sorted(SRC.glob("*.docx"))
    if not files:
        raise SystemExit(f"{SRC} 에 docx 가 없다")

    print(f"{'파일':<36} {'새로':>4} {'덮음':>4} {'표':>4} {'셀':>5}  검증")
    print("-" * 86)
    합_새, 합_덮 = 0, 0
    for p in files:
        doc = Document(str(p))
        before = 통계(doc)
        새로, 덮어씀 = fill(doc)
        합_새 += len(새로)
        합_덮 += len(덮어씀)

        dst = OUT / p.name
        doc.save(str(dst))
        after = 통계(Document(str(dst)))
        표시 = "OK" if before == after else f"구조 {before}→{after}"
        print(
            f"{p.name[:34]:<36} {len(새로):>4} {len(덮어씀):>4} {before[0]:>4} {before[1]:>5}  "
            f"{표시} · {p.stat().st_size:,}→{dst.stat().st_size:,}"
        )

    print(f"\n새로 채운 칸 {합_새}개 · 글자가 있던 칸에 덮어쓴 것 {합_덮}개 · 결과 {OUT}")


if __name__ == "__main__":
    main()
